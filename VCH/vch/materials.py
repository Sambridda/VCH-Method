"""Parses VCH Framework.xlsx into high-resolution smooth curves using PCHIP.
Generates graphs, caches the data for the engine, and checks for updates."""
import os
import json
import shutil
import numpy as np
import openpyxl
from scipy.interpolate import PchipInterpolator

import matplotlib
matplotlib.use('Agg')  # MUST BE BEFORE PYPLOT IMPORT (Thread safety)
import matplotlib.pyplot as plt
from matplotlib.figure import Figure

APP_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
EXCEL_PATH = os.path.join(APP_DIR, "VCH Framework.xlsx")
OUT_DIR = os.path.join(APP_DIR, "out")
MAT_DIR = os.path.join(OUT_DIR, "materials")
CACHE_FILE = os.path.join(OUT_DIR, "materials_cache.json")

FLUIDS = {}
METALS = {}
INSULATION = {}
IMPELLERS = []

def _densify(x, y, num_points=1000):
    if len(x) < 3: return x, y
    spline = PchipInterpolator(x, y)
    x_dense = np.linspace(x.min(), x.max(), num_points)
    return x_dense, spline(x_dense)

def _plot_and_save(name, t, data_cols, labels, ylabel, color, is_fluid=True):
    fig = Figure(figsize=(8, 5))
    ax1 = fig.add_subplot(111)
    
    ax1.plot(t, data_cols[0], color=color, label=labels[0], linewidth=2)
    ax1.set_xlabel("Temperature (°C)")
    ax1.set_ylabel(f"{labels[0]} ({'kg/m³' if is_fluid else 'J/(kg·K)'})", color=color)
    ax1.tick_params(axis='y', labelcolor=color)
    ax1.grid(True, alpha=0.3)
    
    if len(data_cols) > 1:
        ax2 = ax1.twinx()
        ax2.plot(t, data_cols[1], color='tab:red', label=labels[1], linewidth=2, linestyle='--')
        ax2.set_ylabel(f"{labels[1]}", color='tab:red')
        ax2.tick_params(axis='y', labelcolor='tab:red')
        
    ax1.set_title(f"{name} Properties (PCHIP Interpolated)")
    fig.tight_layout()
    safe_name = name.replace(' ', '_').replace('/', '_').lower()
    fig.savefig(os.path.join(MAT_DIR, f"{safe_name}.png"), dpi=100)
    plt.close(fig)

def process_materials(force=False):
    global FLUIDS, METALS, INSULATION, IMPELLERS
    
    if not os.path.exists(EXCEL_PATH):
        raise FileNotFoundError(f"Excel file not found at {EXCEL_PATH}")

    if not force and os.path.exists(CACHE_FILE):
        excel_mtime = os.path.getmtime(EXCEL_PATH)
        cache_mtime = os.path.getmtime(CACHE_FILE)
        if cache_mtime > excel_mtime:
            _load_from_cache()
            return True

    # FIX: IMPELLERS must be a list []
    FLUIDS, METALS, INSULATION, IMPELLERS = {}, {}, {}, []

    # Clear old graphs to ensure we only show the new ones
    if os.path.exists(MAT_DIR):
        shutil.rmtree(MAT_DIR)
    os.makedirs(MAT_DIR, exist_ok=True)

    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    
    # --- 1. FLUIDS ---
    if "Fluids" in wb.sheetnames:
        ws = wb["Fluids"]
        current_name = None
        raw_data = {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0]: current_name = str(row[0]).strip()
            if row[1] is not None and current_name:
                try: t = float(row[1])
                except (ValueError, TypeError): continue
                if np.isnan(t): continue
                
                rho = float(row[3]) if isinstance(row[3], (int, float)) else 1000.0
                cp = float(row[4]) if isinstance(row[4], (int, float)) else 4000.0
                mu = float(row[5]) if isinstance(row[5], (int, float)) else 1.0
                k = float(row[6]) if isinstance(row[6], (int, float)) else 0.5
                
                if current_name not in raw_data: raw_data[current_name] = []
                raw_data[current_name].append((t, rho, cp, mu, k))
        
        for name, data in raw_data.items():
            arr = np.array(sorted(data, key=lambda x: x[0]))
            _, idx = np.unique(arr[:, 0], return_index=True)
            arr = arr[idx]
            t_raw = arr[:, 0]
            
            if len(t_raw) < 3: continue
                
            t_dense, rho_d = _densify(t_raw, arr[:, 1])
            _, cp_d = _densify(t_raw, arr[:, 2])
            _, mu_d = _densify(t_raw, arr[:, 3])
            _, k_d = _densify(t_raw, arr[:, 4])
            
            FLUIDS[name] = np.column_stack((t_dense, rho_d, cp_d, mu_d, k_d)).tolist()
            _plot_and_save(name, t_dense, [rho_d, cp_d], ["Density", "Specific Heat"], "Property", "tab:blue", True)

    # --- 2. METALS ---
    if "Metals" in wb.sheetnames:
        ws = wb["Metals"]
        current_name = None
        raw_data = {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[1]: current_name = str(row[1]).strip()
            if row[2] is not None and current_name:
                try: t = float(row[2])
                except (ValueError, TypeError): continue
                if np.isnan(t): continue
                
                rho = float(row[3]) if isinstance(row[3], (int, float)) else 7800.0
                cp = float(row[4]) if isinstance(row[4], (int, float)) else 500.0
                k = float(row[5]) if isinstance(row[5], (int, float)) else 15.0
                
                if current_name not in raw_data: raw_data[current_name] = []
                raw_data[current_name].append((t, rho, cp, k))
        
        for name, data in raw_data.items():
            arr = np.array(sorted(data, key=lambda x: x[0]))
            _, idx = np.unique(arr[:, 0], return_index=True)
            arr = arr[idx]
            t_raw = arr[:, 0]
            
            if len(t_raw) < 3: continue
                
            t_dense, rho_d = _densify(t_raw, arr[:, 1])
            _, cp_d = _densify(t_raw, arr[:, 2])
            _, k_d = _densify(t_raw, arr[:, 3])
            
            METALS[name] = np.column_stack((t_dense, rho_d, cp_d, k_d)).tolist()
            _plot_and_save(name, t_dense, [cp_d, k_d], ["Specific Heat", "Thermal Cond."], "Property", "tab:green", False)

    # --- 3. INSULATION & IMPELLERS ---
    if "Insulation" in wb.sheetnames:
        ws = wb["Insulation"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0] and row[1]:
                try: INSULATION[str(row[0]).strip()] = float(row[1])
                except (ValueError, TypeError): pass

    if "Impellers" in wb.sheetnames:
        ws = wb["Impellers"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0] and row[2] and row[3]:
                try:
                    IMPELLERS.append(dict(name=str(row[0]).strip(), pattern=str(row[1]).strip() if row[1] else "",
                                          C=float(row[2]), a=float(row[3])))
                except (ValueError, TypeError): pass

    # --- SAVE CACHE ---
    cache_data = dict(FLUIDS=FLUIDS, METALS=METALS, INSULATION=INSULATION, IMPELLERS=IMPELLERS)
    with open(CACHE_FILE, "w") as f:
        json.dump(cache_data, f)
        
    return True

def _load_from_cache():
    global FLUIDS, METALS, INSULATION, IMPELLERS
    if not os.path.exists(CACHE_FILE): return
    with open(CACHE_FILE, "r") as f:
        data = json.load(f)
    
    FLUIDS = {k: np.array(v) for k, v in data.get("FLUIDS", {}).items()}
    METALS = {k: np.array(v) for k, v in data.get("METALS", {}).items()}
    INSULATION = data.get("INSULATION", {})
    IMPELLERS = data.get("IMPELLERS", [])

if os.path.exists(CACHE_FILE):
    _load_from_cache()